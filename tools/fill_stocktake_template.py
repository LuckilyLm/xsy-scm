#!/usr/bin/env python3
"""把盘点快照导入模板填成「可成功导入」的实盘件（仅供 E2E / 手工验收使用）。

为什么需要：Wave 6 的浏览器验收要求真实跑一遍「导出模板 → 填实盘 → 导入建草稿」，
而前端页面本身不解析 xlsx（导入是整文件上传给后端），所以“填好的文件”必须由测试侧生成。

做法与约束（对齐 `InventoryStocktakeImportService` 的读法）：
  · 只写「实盘数量」列（第 7 列），快照凭证 / SKU 编码 / 账面快照等受保护列**逐字节不动**——
    凭证是 HMAC 签名的字符串单元格，重写它等于伪造来源集合；
  · 填的是「账面快照 + delta」，确保确认阶段真的产生盘盈/盘亏流水（漂移反例依赖这一点）；
  · 以字符串写入：后端正则 `[0-9]{1,14}(\\.[0-9]{1,4})?` 按文本校验，数值单元格会经
    `NumberToTextConverter` 转换，留成字符串可避免浮点尾巴；
  · 结果为负时夹到 0（实盘量不允许负数）。

用法：
    python tools/fill_stocktake_template.py --in template.xlsx --out filled.xlsx [--delta 1]
输出（stdout，单行 JSON）：
    {"rows": 3, "actuals": {"SKU-A": "12.5", ...}}
"""

import argparse
import json
import sys
from decimal import Decimal, InvalidOperation

ACTUAL_COLUMN = 7      # 1-based，对应表头「实盘数量」
BOOK_COLUMN = 6        # 1-based，对应表头「账面数量快照」


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--in", dest="src", required=True)
    parser.add_argument("--out", dest="dst", required=True)
    parser.add_argument("--delta", default="1")
    args = parser.parse_args()

    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError:
        print("[fill-stocktake] 缺少 python 包 openpyxl。请执行：\n"
              "  python -m pip install openpyxl", file=sys.stderr)
        return 2

    try:
        delta = Decimal(args.delta)
    except InvalidOperation:
        print(f"[fill-stocktake] --delta 不是合法数字: {args.delta!r}", file=sys.stderr)
        return 2

    # data_only=False：必须保留公式与全部原始单元格的值，凭证列不能被“计算结果”替换。
    book = load_workbook(args.src)
    sheet = book.active
    actuals: dict[str, str] = {}
    for row in range(2, sheet.max_row + 1):
        sku_code = sheet.cell(row=row, column=2).value
        if sku_code is None or str(sku_code).strip() == "":
            continue  # 空行：模板末尾可能有装饰性空行，不能给它填值
        snapshot = str(sheet.cell(row=row, column=BOOK_COLUMN).value or "0").strip()
        try:
            value = Decimal(snapshot) + delta
        except InvalidOperation:
            print(f"[fill-stocktake] 第 {row} 行账面快照不是数字: {snapshot!r}", file=sys.stderr)
            return 2
        if value < 0:
            value = Decimal(0)
        text = format(value.normalize(), "f")
        sheet.cell(row=row, column=ACTUAL_COLUMN).value = text
        actuals[str(sku_code).strip()] = text

    if not actuals:
        print("[fill-stocktake] 模板没有可填的数据行（该仓库可能没有库存余额）", file=sys.stderr)
        return 2

    book.save(args.dst)
    print(json.dumps({"rows": len(actuals), "actuals": actuals}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
